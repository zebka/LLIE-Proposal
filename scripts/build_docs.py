# -*- coding: utf-8 -*-
"""Build TWO deliverables from the same md sources:
  1) Proposal-Standalone.docx  — standalone Word (no sample form), figures embedded
  2) latex/main.tex            — XePersian LaTeX version (same figures)
Run:  python build_docs.py
"""
import io
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, str(Path(__file__).parent))
from fill_form import (parse_bib, ieee_ref, clean_md, build_cite_map,
                       apply_cites, map_cites, TABLES, RISK_ROWS,
                       TIMELINE_ROWS, COMPARISON_ROWS, MILESTONE, PROJ)

BASE = PROJ.parent
SRC = PROJ / "src"
FIGS = PROJ / "figures"
LATEX = PROJ / "latex"
OUTPUT = PROJ / "output"
DOCX_OUT = OUTPUT / "Proposal-Standalone.docx"

FA_TITLE = ("ارائه یک روش بهبود تصویر کم‌نور مبتنی بر پیشین (Prior) مدل‌های انتشار (Diffusion) "
            "برای شرایط نوری نامتوازن گرگ‌ومیش با کاربرد در بینایی ماشین")
EN_TITLE = ("A Diffusion-Prior Zero-Shot Method for Low-Light Image Enhancement under "
            "Imbalanced Twilight Conditions with Application to Machine Vision")
FA_KEYS = "بهبود تصویر کم‌نور، مدل‌های انتشار، یادگیری بدون نظارت، گرگ‌ومیش، بینایی ماشین"
EN_KEYS = "Low-Light Image Enhancement, Diffusion Models, Zero-Shot Learning, Twilight, Machine Vision"

TABLES = dict(TABLES)
TABLES["5"] = TIMELINE_ROWS
TABLES["COMP"] = COMPARISON_ROWS

FIG_FILES = {n: f"شکل-{n:02d}.png" for n in range(1, 15)}
FIG_FILES[15] = ["شکل-16a.png", "شکل-16b.png", "شکل-16c.png"]
FIG_FILES[16] = None  # user's own photos (phase 3)
FIG_W = {1: 8.5, 4: 10.0, 5: 10.0, 11: 8.5, 12: 8.5}  # cm; default 15.5

SEC_TITLES = [
    ("۱", "شرح مسئله (اهداف، سابقه و ضرورت تحقیق)"),
    ("۲", "سوالات تحقیق/فرضیه‌ها"),
    ("۳", "کاربردهای تحقیق و استفاده‌کنندگان از نتایج رساله"),
    ("۴", "روش شناسی (روش و ابزار جمع‌آوری، انجام و تحلیل اطلاعات تحقیق)"),
    ("۵", "جنبه نوآوری تحقیق"),
    ("۶", "جدول زمان‌بندی انجام پروژه"),
]

FIG_RE = re.compile(r"^شکل ([۰-۹]+):")
PD = "۰۱۲۳۴۵۶۷۸۹"
fa2en = lambda s: int(s.translate(str.maketrans(PD, "0123456789")))


# ---------------------------------------------------------------- docx helpers
def set_rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)


def add_para(doc, text, rtl=True, bold=False, size=11, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else (
        WD_ALIGN_PARAGRAPH.JUSTIFY if rtl else WD_ALIGN_PARAGRAPH.LEFT)
    if rtl:
        set_rtl(p)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p


def add_table(doc, rows, size=9):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or ci == 0 and len(rows[0]) <= 4) else WD_ALIGN_PARAGRAPH.JUSTIFY
            set_rtl(p)
            run = p.add_run(val)
            run.bold = ri == 0
            run.font.size = Pt(size)
    return t


def add_figure(doc, fig_num, files, width_default=15.5):
    for f in files:
        if (FIGS / f).exists():
            w = FIG_W.get(fig_num, width_default)
            doc.add_picture(str(FIGS / f), width=Cm(w))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


# ---------------------------------------------------------------- latex helpers
def esc(s):
    for ch, rep in [("\\", r"\textbackslash{}"), ("&", r"\&"), ("%", r"\%"),
                    ("#", r"\#"), ("_", r"\_"), ("{", r"\{"), ("}", r"\}"),
                    ("~", r"\textasciitilde{}"), ("^", r"\textasciicircum{}")]:
        s = s.replace(ch, rep)
    return s


LAT_RAW = re.compile(r"[A-Za-z0-9][A-Za-z0-9+\-./()%@~–]*")


def esc_tex(s):
    """Escape LaTeX specials + wrap Latin/numeric runs in \\lr{} so they
    get the Latin font (B Nazanin Bold has no Latin glyphs) and keep LTR order."""
    s = s.replace("—", "ـــ")  # em-dash not in B Nazanin; use tatweel
    out = []
    pos = 0
    for m in LAT_RAW.finditer(s):
        out.append(esc(s[pos:m.start()]).replace("–", "ـــ"))
        out.append("\\lr{" + esc(m.group(0)) + "}")
        pos = m.end()
    out.append(esc(s[pos:]).replace("–", "ـــ"))
    r = "".join(out)
    r = (r.replace("٪", "\\%").replace("•", "ــ")
          .replace("<", "$<$").replace(">", "$>$")
          .replace("≤", "$\\leq$").replace("≥", "$\\geq$")
          .replace("≈", "$\\approx$").replace("±", "$\\pm$")
          .replace("×", "$\\times$")
          .replace("✅", "$\\checkmark$").replace("❌", "$\\times$")
          .replace("\\textasciitilde{}", "$\\sim$")
          .replace("٬", ","))
    return r


def latex_table(rows, small=True):
    ncol = len(rows[0])
    colspec = "|c|" + "c|" * (ncol - 1)
    size = "\\small" if small else ""
    lines = [f"{{{size}", f"\\begin{{tabular}}{{{colspec}}}", "\\hline"]
    for ri, row in enumerate(rows):
        cells = " & ".join(esc_tex(c) for c in row)
        lines.append(cells + " \\\\ \\hline")
    lines.append("\\end{tabular}}")
    return "\n".join(lines)


def main():
    entries = parse_bib(PROJ / "references.bib")
    sec1 = clean_md((SRC / "section1-draft.md").read_text(encoding="utf-8").split("## ۱-۱.", 1)[1].join(["## ۱-۱.", ""]))
    all_md = (SRC / "sections-2-3-4-draft.md").read_text(encoding="utf-8")
    p2 = all_md.split("# (۲)", 1)[1]
    sec2 = clean_md(p2.split("# (۳)", 1)[0])
    sec3 = clean_md(p2.split("# (۳)", 1)[1].split("# (۴)", 1)[0])
    sec4 = clean_md(p2.split("# (۴)", 1)[1])
    md56 = (SRC / "sections-5-6-draft.md").read_text(encoding="utf-8")
    sec5 = clean_md(md56.split("# (۵)", 1)[1].split("# (۶)", 1)[0])
    sec6 = [("h", "برنامه زمان‌بندی"),
            ("p", "زمان‌بندی اجرایی پروژه در قالب پنج فاز طی ۲۴ ماه مطابق جدول زیر است."),
            ("p", "جدول ۵: برنامه زمان‌بندی پنج‌فازی اجرای پروژه (تألیف پژوهش حاضر)"),
            ("table", "5"),
            ("p", MILESTONE)]

    order = build_cite_map([sec1, sec2, sec3, sec4, sec5, sec6], entries)
    sec1, sec2, sec3, sec4, sec5, sec6 = [apply_cites(s, order) for s in (sec1, sec2, sec3, sec4, sec5, sec6)]
    for key, rows in TABLES.items():
        for i, row in enumerate(rows):
            rows[i] = [map_cites(v, order) for v in row]
    sections = [(num, title, items) for (num, title), items in
                zip(SEC_TITLES, [sec1, sec2, sec3, sec4, sec5, sec6])]

    # ============================================================ DOCX
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.3)
        s.left_margin = s.right_margin = Cm(2.2)

    add_para(doc, FA_TITLE, bold=True, size=14, center=True)
    add_para(doc, EN_TITLE, rtl=False, size=11, center=True)
    add_para(doc, "کلیدواژه‌ها: " + FA_KEYS, size=10.5)
    add_para(doc, "Keywords: " + EN_KEYS, rtl=False, size=10)

    info_rows = [
        ["نام و نام خانوادگی دانشجو", "............................", "شماره دانشجویی", "............................"],
        ["دانشکده / رشته / گرایش", "مهندسی برق — الکترونیک", "نوع پذیرش", "............................"],
        ["استاد راهنما", "............................", "استاد مشاور", "............................"],
        ["مدت اجرا", "۲۴ ماه", "تعداد واحد", "۲۱"],
    ]
    add_table(doc, info_rows, size=9.5)

    for num, title, items in sections:
        add_para(doc, f"{num}) {title}", bold=True, size=13)
        for kind, txt in items:
            if kind == "sp":
                continue
            if kind == "table":
                add_table(doc, TABLES[txt])
                continue
            m = FIG_RE.match(txt)
            if m and kind == "p":
                fnum = fa2en(m.group(1))
                files = FIG_FILES.get(fnum)
                if files:
                    add_figure(doc, fnum, files if isinstance(files, list) else [files])
                else:
                    add_para(doc, "[ جای تصاویر گرگ‌ومیش ضبط‌شده در فاز ۳، دیتاست TwilightDrive ]",
                             size=10, center=True)
            add_para(doc, txt, bold=(kind == "h"),
                     size=12 if kind == "h" else 10.5)

    add_para(doc, "۷) مراجع", bold=True, size=13)
    for key, num in sorted(order.items(), key=lambda kv: kv[1]):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(f"[{num}] {ieee_ref(key, entries)}").font.size = Pt(9)

    OUTPUT.mkdir(parents=True, exist_ok=True)
    doc.save(str(DOCX_OUT))

    # ============================================================ LaTeX
    lf = LATEX / "figures"
    lf.mkdir(parents=True, exist_ok=True)
    tex_figs = {}
    for n, files in FIG_FILES.items():
        if files is None:
            continue
        files = files if isinstance(files, list) else [files]
        out = []
        for i, f in enumerate(files):
            tag = f"fig-{n:02d}" + (chr(96 + i) if len(files) > 1 else "") + ".png"
            shutil.copy(FIGS / f, lf / tag)
            out.append(tag)
        tex_figs[n] = out

    L = []
    L.append("% !TeX program = xelatex")
    L.append("% Compile:  xelatex main.tex   (run twice)   — requires the xepersian package")
    L.append("% Fonts: B Nazanin (standard in Iran). Alternatives: XB Zar, XB Niloofar")
    L.append("\\documentclass[12pt,a4paper]{article}")
    L.append("\\usepackage[margin=2.3cm]{geometry}")
    L.append("\\usepackage{amsmath,amssymb,graphicx,float,booktabs,longtable}")
    L.append("\\usepackage{xepersian}")
    L.append("\\settextfont{B Nazanin}")
    L.append("\\setlatintextfont[BoldFont={Times New Roman}, ItalicFont={Times New Roman}, BoldItalicFont={Times New Roman}]{Times New Roman}")
    L.append("\\begin{document}")
    L.append("")
    L.append("\\begin{center}")
    L.append("{\\Large \\textbf{" + esc_tex(FA_TITLE) + "}}\\\\[6pt]")
    L.append("{\\normalsize \\lr{" + esc(EN_TITLE) + "}}")
    L.append("\\end{center}")
    L.append("\\noindent\\textbf{کلیدواژه‌ها:} " + esc_tex(FA_KEYS) + "\\\\")
    L.append("\\noindent\\textbf{\\lr{Keywords:}} \\lr{" + esc(EN_KEYS) + "}")
    L.append("")
    for num, title, items in sections:
        L.append("")
        L.append("\\section{" + esc_tex(title) + "}")
        for kind, txt in items:
            if kind == "sp":
                continue
            if kind == "table":
                L.append("\\begin{table}[H]\\centering")
                L.append(latex_table(TABLES[txt]))
                L.append("\\end{table}")
                continue
            m = FIG_RE.match(txt)
            if m and kind == "p":
                fnum = fa2en(m.group(1))
                files = tex_figs.get(fnum)
                if files:
                    L.append("\\begin{figure}[H]\\centering")
                    for f in files:
                        L.append(f"\\includegraphics[width=0.92\\linewidth]{{figures/{f}}}\\\\[4pt]")
                    L.append("\\end{figure}")
                else:
                    L.append("\\begin{center}\\fbox{جای تصاویر گرگ‌ومیش فاز ۳ ـــ \\lr{TwilightDrive}}\\end{center}")
            body = esc_tex(txt)
            if kind == "h":
                L.append("\\textbf{" + body + "}\\par")
            else:
                L.append(body + "\\par")
        L.append("")
    L.append("")
    L.append("\\section{مراجع}")
    L.append("\\begin{thebibliography}{" + str(len(order)) + "}")
    for key, num in sorted(order.items(), key=lambda kv: kv[1]):
        L.append(f"\\bibitem{{{key}}} \\lr{{" + esc(ieee_ref(key, entries)) + "}")
    L.append("\\end{thebibliography}")
    L.append("\\end{document}")
    (LATEX / "main.tex").write_text("\n".join(L), encoding="utf-8")

    return order


if __name__ == "__main__":
    order = main()
    print(f"OK — DOCX: {DOCX_OUT.name}")
    print(f"OK — LaTeX: {LATEX / 'main.tex'} ({len(order)} bibitems)")
